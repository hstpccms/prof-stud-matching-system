"""
ตัวสร้างไฟล์ Excel สำหรับ "ดาวน์โหลด" (แยกจากไฟล์ที่ feed อัลกอริทึม matching)

สำคัญ: โมดูลนี้ไม่เกี่ยวกับการ run matching เลย — อัลกอริทึมยังอ่านไฟล์จาก
`_export_input_excel` (schema เดิม) และ `_import_results` อ่านผลจากไฟล์ native เดิม
ที่นี่แค่ "เรนเดอร์ใหม่เพื่อแสดง/ดาวน์โหลด" ให้สวยและอ่านง่ายเท่านั้น
"""
import json
import re
from typing import Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

import models

_HEADER_FONT = Font(bold=True, size=11, color="1F2937")
_HEADER_FILL = PatternFill("solid", fgColor="E6F4FF")
_TITLE_FONT = Font(bold=True, size=13, color="1F4E79")
_TITLE_FILL = PatternFill("solid", fgColor="D9E2F3")
_ZEBRA_FILL = PatternFill("solid", fgColor="F2F6FB")
_ADV_FONT = Font(bold=True, color="217346")
_BORDER = Border(
    left=Side(style="thin", color="D1D5DB"), right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"), bottom=Side(style="thin", color="D1D5DB"),
)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _natkey(code):
    s = str(code or "")
    m = re.findall(r"\d+", s)
    return (int(m[0]) if m else 0, s)


def _style_header(ws, ncols: int, row: int = 1):
    ws.row_dimensions[row].height = 26
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _BORDER


def _autosize(ws, min_widths=None):
    min_widths = min_widths or {}
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        longest = 0
        for cell in col:
            for line in str(cell.value or "").split("\n"):
                longest = max(longest, len(line))
        ws.column_dimensions[letter].width = min(max(longest + 4, min_widths.get(letter, 12)), 55)


# ════════════════════════════════════════════════════════════════════════════
# RAW DATA (display) — Group_Info / Professor_Info / Student_Rankings / Professor_Scores
# ════════════════════════════════════════════════════════════════════════════
def export_rawdata_display(session_id: int, program: Optional[str], db: Session, out_path: str):
    q_groups = db.query(models.Group).filter_by(session_id=session_id)
    q_profs = db.query(models.Professor).filter_by(session_id=session_id)
    q_rankings = db.query(models.StudentRanking).filter_by(session_id=session_id)
    q_scores = db.query(models.ProfessorScore).filter_by(session_id=session_id)
    if program:
        has_prog = db.query(models.Group).filter(
            models.Group.session_id == session_id, models.Group.program != None
        ).count() > 0
        if has_prog:
            q_groups = q_groups.filter_by(program=program)
            q_profs = q_profs.filter_by(program=program)
            q_rankings = q_rankings.filter_by(program=program)
            q_scores = q_scores.filter_by(program=program)

    groups = sorted(q_groups.all(), key=lambda g: _natkey(g.anonymous_code or g.group_id or g.id))
    professors = sorted(q_profs.all(), key=lambda p: _natkey(p.anonymous_code or p.prof_id or p.id))
    rankings = q_rankings.all()
    scores = q_scores.all()

    # map รหัสอาจารย์ -> ชื่อ (ไว้ใช้แทนรหัสในชีท Student_Rankings / Professor_Scores)
    prof_codes = [p.anonymous_code or p.prof_id for p in professors]
    prof_name_of = {(p.anonymous_code or p.prof_id): (p.full_name or (p.anonymous_code or p.prof_id)) for p in professors}

    wb = openpyxl.Workbook()

    # ── 1. Group_Info ──
    ws = wb.active
    ws.title = "Group_Info"
    headers = ["GroupID", "จำนวนสมาชิก", "รหัสนักศึกษา", "ชื่อ-นามสกุล", "หัวข้อที่สนใจ", "รายละเอียดของหัวข้อ"]
    ws.append(headers)
    _style_header(ws, len(headers))
    for ridx, g in enumerate(groups, start=2):
        gid = g.anonymous_code or g.group_id or f"G{ridx-1:03d}"
        ids, names = [], []
        for m in (g.members or []):
            if m.student_id:
                ids.append(str(m.student_id).strip())
            if m.full_name:
                names.append(str(m.full_name).strip())
        try:
            topics = json.loads(g.topic_interest) if g.topic_interest else []
        except Exception:
            topics = []
        titles, details = [], []
        for i, t in enumerate(topics, start=1) if isinstance(topics, list) else []:
            if isinstance(t, dict):
                title, detail = (t.get("title") or "").strip(), (t.get("detail") or "").strip()
            else:
                title, detail = str(t).strip(), ""
            if title:
                titles.append(f"{i}. {title}")
                details.append(f"{i}. {detail if detail and detail != '—' else '—'}")
        count = g.member_count or (len(g.members) if g.members else len(ids)) or 1
        row = [gid, count, "\n".join(ids) or "—", "\n".join(names) or "—",
               "\n".join(titles) or "—", "\n".join(details) or "—"]
        ws.append(row)
        for c in range(1, len(row) + 1):
            cell = ws.cell(row=ridx, column=c)
            cell.border = _BORDER
            cell.alignment = _CENTER if c in (1, 2) else _LEFT
    _autosize(ws, {"A": 12, "B": 14, "C": 18, "D": 24, "E": 34, "F": 42})

    # ── 2. Professor_Info ──
    ws2 = wb.create_sheet("Professor_Info")
    headers = ["ProfID", "ชื่อ-นามสกุล", "ความเชี่ยวชาญ", "โควต้ากลุ่ม"]
    ws2.append(headers)
    _style_header(ws2, len(headers))
    for ridx, p in enumerate(professors, start=2):
        row = [p.anonymous_code or p.prof_id or f"P{ridx-1:03d}", p.full_name or "—",
               p.expertise or "—", p.quota if p.quota is not None else 0]
        ws2.append(row)
        for c in range(1, len(row) + 1):
            cell = ws2.cell(row=ridx, column=c)
            cell.border = _BORDER
            cell.alignment = _CENTER if c in (1, 4) else _LEFT
    _autosize(ws2, {"A": 12, "B": 24, "C": 40, "D": 12})

    # ── 3. Student_Rankings (หัวคอลัมน์ = ชื่ออาจารย์) ──
    ws3 = wb.create_sheet("Student_Rankings")
    headers = ["GroupID"] + [prof_name_of.get(pc, pc) for pc in prof_codes]
    ws3.append(headers)
    _style_header(ws3, len(headers))
    rank_lookup = {(r.group_code, r.prof_code): r.rank for r in rankings}
    for ridx, g in enumerate(groups, start=2):
        gcode = g.anonymous_code or g.group_id
        row = [gcode] + [rank_lookup.get((gcode, pc)) for pc in prof_codes]
        ws3.append(row)
        for c in range(1, len(row) + 1):
            cell = ws3.cell(row=ridx, column=c)
            cell.border = _BORDER
            cell.alignment = _CENTER
    _autosize(ws3, {"A": 14})

    # ── 4. Professor_Scores (ชื่ออาจารย์แทนรหัส, ไม่มี MainScore, SubScore→Score) ──
    ws4 = wb.create_sheet("Professor_Scores")
    headers = ["ชื่ออาจารย์", "GroupID", "Score_TopicFit_A", "Score_Clarity_B", "Score"]
    ws4.append(headers)
    _style_header(ws4, len(headers))
    sorted_scores = sorted(scores, key=lambda s: (_natkey(s.prof_code), _natkey(s.group_code)))
    for ridx, s in enumerate(sorted_scores, start=2):
        row = [prof_name_of.get(s.prof_code, s.prof_code), s.group_code, s.score_a, s.score_b, s.sub_score]
        ws4.append(row)
        for c in range(1, len(row) + 1):
            cell = ws4.cell(row=ridx, column=c)
            cell.border = _BORDER
            cell.alignment = _LEFT if c == 1 else _CENTER
    _autosize(ws4, {"A": 24, "B": 12})

    wb.save(out_path)


# ════════════════════════════════════════════════════════════════════════════
# RESULT (display) — สร้างจาก DB + อ่าน Tie-Break จากไฟล์ native เดิม
# ════════════════════════════════════════════════════════════════════════════
def _mode_stats(rows):
    n = len(rows)
    matched = [r for r in rows if r.assigned_prof and r.assigned_prof != "UNMATCHED"]
    ranks = [r.rank_given for r in matched if r.rank_given is not None]
    subs = [r.sub_score for r in matched if r.sub_score is not None]
    return {
        "n_matched": len(matched),
        "n_unmatched": n - len(matched),
        "avg_rank": round(sum(ranks) / len(ranks), 2) if ranks else None,
        "pct_top1": round(100 * sum(1 for x in ranks if x == 1) / n, 1) if n else None,
        "pct_top3": round(100 * sum(1 for x in ranks if x <= 3) / n, 1) if n else None,
        "avg_score": round(sum(subs) / len(subs), 2) if subs else None,
    }


def _write_matching_sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    headers = ["GroupCode", "AssignedProfessor", "RankGroupGaveProf(1=best)", "ScoreProfGaveGroup"]
    ws.append(headers)
    _style_header(ws, len(headers))
    for ridx, r in enumerate(sorted(rows, key=lambda x: _natkey(x.group_code)), start=2):
        row = [r.group_code, r.assigned_prof or "UNMATCHED", r.rank_given, r.sub_score]
        ws.append(row)
        for c in range(1, len(row) + 1):
            cell = ws.cell(row=ridx, column=c)
            cell.border = _BORDER
            cell.alignment = _CENTER
    _autosize(ws, {"A": 12, "B": 18, "C": 24, "D": 18})
    ws.freeze_panes = "A2"


def _write_prof_summary_sheet(wb, title, rows, quota_of):
    ws = wb.create_sheet(title)
    headers = ["ProfCode", "Quota", "GroupsAssigned", "NumAssigned", "QuotaRemaining"]
    ws.append(headers)
    _style_header(ws, len(headers))
    assigned = {}
    for r in rows:
        if r.assigned_prof and r.assigned_prof != "UNMATCHED":
            assigned.setdefault(r.assigned_prof, []).append(r.group_code)
    ridx = 2
    for code in sorted(quota_of, key=_natkey):
        gs = sorted(assigned.get(code, []), key=_natkey)
        quota = quota_of.get(code, 0) or 0
        row = [code, quota, ", ".join(gs) if gs else "-", len(gs), quota - len(gs)]
        ws.append(row)
        for c in range(1, len(row) + 1):
            cell = ws.cell(row=ridx, column=c)
            cell.border = _BORDER
            cell.alignment = _LEFT if c == 3 else _CENTER
        ridx += 1
    _autosize(ws, {"A": 12, "B": 8, "C": 40, "D": 14, "E": 16})


def _write_comparison_sheet(wb, st_student, st_prof, seed):
    ws = wb.create_sheet("Comparison")

    def better(a, b, higher=True):
        if a is None or b is None:
            return "—"
        if a == b:
            return "เท่ากัน"
        if higher:
            return "ฝั่งนักศึกษา" if a > b else "ฝั่งอาจารย์"
        return "ฝั่งนักศึกษา" if a < b else "ฝั่งอาจารย์"

    # Title row
    ws.append(["สรุปเปรียบเทียบผลการจับคู่ 2 รูปแบบ"])
    ws.merge_cells("A1:D1")
    tcell = ws.cell(row=1, column=1)
    tcell.font = _TITLE_FONT
    tcell.fill = _TITLE_FILL
    tcell.alignment = _CENTER
    ws.row_dimensions[1].height = 30

    # Header row
    headers = ["ตัวชี้วัด", "แบบนักศึกษาเสนอ", "แบบอาจารย์เสนอ", "ฝั่งที่ได้เปรียบ"]
    ws.append(headers)
    _style_header(ws, len(headers), row=2)

    rows = [
        ("จำนวนกลุ่มที่จับคู่สำเร็จ", st_student["n_matched"], st_prof["n_matched"],
         better(st_student["n_matched"], st_prof["n_matched"])),
        ("จำนวนกลุ่มที่ไม่ได้จับคู่", st_student["n_unmatched"], st_prof["n_unmatched"],
         better(st_student["n_unmatched"], st_prof["n_unmatched"], higher=False)),
        ("อันดับเฉลี่ยที่นักศึกษาได้รับ (ยิ่งน้อยยิ่งดี)", st_student["avg_rank"], st_prof["avg_rank"],
         better(st_student["avg_rank"], st_prof["avg_rank"], higher=False)),
        ("% กลุ่มที่ได้อาจารย์อันดับ 1 ของตน", st_student["pct_top1"], st_prof["pct_top1"],
         better(st_student["pct_top1"], st_prof["pct_top1"])),
        ("% กลุ่มที่ได้อาจารย์ใน Top 3 ของตน", st_student["pct_top3"], st_prof["pct_top3"],
         better(st_student["pct_top3"], st_prof["pct_top3"])),
        ("คะแนนเฉลี่ยที่อาจารย์ให้กลุ่มที่ได้รับ (ความพอใจฝั่งอาจารย์)", st_student["avg_score"], st_prof["avg_score"],
         better(st_student["avg_score"], st_prof["avg_score"])),
    ]
    r0 = 3
    for i, row in enumerate(rows):
        ws.append(list(row))
        rownum = r0 + i
        for c in range(1, 5):
            cell = ws.cell(row=rownum, column=c)
            cell.border = _BORDER
            cell.alignment = _LEFT if c == 1 else _CENTER
            if i % 2 == 1:
                cell.fill = _ZEBRA_FILL
            if c == 4:
                cell.font = _ADV_FONT
    # Seed footer
    seed_row = r0 + len(rows) + 1
    ws.cell(row=seed_row, column=1, value="Seed สำหรับตัดสินเสมอ (ตรวจสอบย้อนหลัง)")
    ws.cell(row=seed_row, column=2, value=seed)
    ws.cell(row=seed_row, column=1).font = Font(italic=True, color="666666")
    ws.cell(row=seed_row, column=2).font = Font(bold=True)
    _autosize(ws, {"A": 46, "B": 18, "C": 18, "D": 16})


def _write_tiebreak_sheet(wb, native_path):
    ws = wb.create_sheet("Tie-Break Log")
    headers = ["ProfCode", "GroupsTied", "Score", "RankByGroup", "ResolvedOrder(bestFirst)"]
    ws.append(headers)
    _style_header(ws, len(headers))

    data_rows = []
    if native_path:
        try:
            src = openpyxl.load_workbook(native_path, data_only=True)
            if "TieBreak_Log" in src.sheetnames:
                sws = src["TieBreak_Log"]
                for row in sws.iter_rows(min_row=2, values_only=True):
                    if not row or all(v is None or str(v).strip() == "" for v in row):
                        continue
                    first = str(row[0] or "").strip()
                    if first.startswith("หมายเหตุ"):   # ตัดหมายเหตุออก (ข้อ 3.4)
                        continue
                    data_rows.append(list(row[:5]) + [None] * (5 - len(row[:5])))
        except Exception:
            pass

    if not data_rows:
        data_rows = [["-", "ไม่มี Tie ที่ต้องใช้ Seeded Random ในรอบนี้", "", "", ""]]

    for ridx, row in enumerate(data_rows, start=2):
        ws.append(row)
        for c in range(1, 6):
            cell = ws.cell(row=ridx, column=c)
            cell.border = _BORDER
            cell.alignment = _CENTER if c in (1, 3, 4) else _LEFT
    _autosize(ws, {"A": 12, "B": 24, "C": 10, "D": 14, "E": 32})


def export_result_display(run_id: int, db: Session, out_path: str):
    run = db.query(models.MatchingRun).filter_by(id=run_id).first()
    if not run:
        raise ValueError("ไม่พบ run")

    q_prof = db.query(models.Professor).filter_by(session_id=run.session_id)
    if run.program:
        q_prof = q_prof.filter_by(program=run.program)
    quota_of = {(p.anonymous_code or p.prof_id): (p.quota or 0) for p in q_prof.all() if (p.anonymous_code or p.prof_id)}

    all_results = db.query(models.MatchingResult).filter_by(run_id=run_id).all()
    student_rows = [r for r in all_results if r.mode == "student"]
    prof_rows = [r for r in all_results if r.mode == "professor"]
    if not student_rows and not prof_rows:
        student_rows = all_results  # legacy single-mode

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    if student_rows:
        _write_matching_sheet(wb, "Matching (Student-Proposing)", student_rows)
    if prof_rows:
        _write_matching_sheet(wb, "Matching (Professor-Proposing)", prof_rows)
    if student_rows:
        _write_prof_summary_sheet(wb, "Professor Summary (Student)", student_rows, quota_of)
    if prof_rows:
        _write_prof_summary_sheet(wb, "Professor Summary (Professor)", prof_rows, quota_of)

    if student_rows and prof_rows:
        _write_comparison_sheet(wb, _mode_stats(student_rows), _mode_stats(prof_rows), run.seed)

    _write_tiebreak_sheet(wb, run.output_file_path)

    wb.save(out_path)
